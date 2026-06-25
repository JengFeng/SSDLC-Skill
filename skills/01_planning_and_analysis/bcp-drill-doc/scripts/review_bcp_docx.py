# -*- coding: utf-8 -*-
"""
BCP docx 靜態檢查 — 掃常見退件點，按 🔴🟡🟢 分級

用法：
    python review_bcp_docx.py --input existing.docx
"""
import argparse
import re
from docx import Document


# === 規則集 ===
# Each rule returns list of tuples: (severity, location, message)
#   severity in {'RED', 'YELLOW', 'GREEN'}


def rule_unit_misplacement(doc):
    findings = []
    forbidden = ['資訊室', '資料治理科', '氣象署']
    for i, p in enumerate(doc.paragraphs):
        for kw in forbidden:
            if kw in p.text:
                findings.append((
                    'RED',
                    f'Paragraph [{i}]',
                    f'出現誤植單位「{kw}」 → 應為「資訊管理科」'
                ))
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for kw in forbidden:
                    if kw in cell.text:
                        findings.append((
                            'RED',
                            f'Table {ti} R{ri}C{ci}',
                            f'出現誤植單位「{kw}」 → 應為「資訊管理科」'
                        ))
    return findings


def rule_metrics_missing(doc):
    findings = []
    full_text = '\n'.join(p.text for p in doc.paragraphs)
    if 'MTPD' not in full_text:
        findings.append(('RED', '伍/柒 時間管控指標', 'MTPD 完全沒寫'))
    if 'RTO' not in full_text:
        findings.append(('RED', '伍/柒 時間管控指標', 'RTO 完全沒寫'))
    if 'RPO' not in full_text:
        findings.append(('RED', '伍/柒 時間管控指標', 'RPO 完全沒寫'))

    # check for "依實際作業" / "本項作業需求完成時間：" 空白後接非數字
    for i, p in enumerate(doc.paragraphs):
        if '本項作業需求完成時間' in p.text:
            after = p.text.split('：', 1)[-1].strip() if '：' in p.text else ''
            if not after or after in ('依實際作業', '(依實際作業 )', '(依實際作業)'):
                findings.append((
                    'RED',
                    f'Paragraph [{i}]',
                    f'完成時間欄寫「{after or "空白"}」 → 請填具體分鐘/小時數'
                ))
    return findings


def rule_duplicate_phrase_in_purpose(doc):
    findings = []
    for i, p in enumerate(doc.paragraphs[:60]):
        if '實施備份還原演練' in p.text and p.text.count('並實施備份還原演練') >= 2:
            findings.append((
                'YELLOW',
                f'壹、目的 [{i}]',
                '「並實施備份還原演練」重複出現'
            ))
        if '。。' in p.text:
            findings.append((
                'YELLOW',
                f'Paragraph [{i}]',
                '兩個句號「。。」連在一起'
            ))
    return findings


def rule_chinese_terms_typo(doc):
    findings = []
    mainland_terms = {
        '技術棧': '技術環境／平台規格',
        '代碼': '程式碼',
        '服務器': '伺服器',
        '數據庫': '資料庫',
        '數據': '資料',
        '默認': '預設',
        '軟件': '軟體',
        '視頻': '影片',
        '打印': '列印',
        '硬盤': '硬碟',
        '內存': '記憶體',
        '文件夾': '資料夾',
        '賬號': '帳號',
    }
    for i, p in enumerate(doc.paragraphs):
        for k, v in mainland_terms.items():
            if k in p.text:
                findings.append((
                    'YELLOW',
                    f'Paragraph [{i}]',
                    f'大陸用語「{k}」 → 建議改「{v}」'
                ))
    return findings


def rule_empty_destruction(doc):
    findings = []
    if not doc.tables:
        return findings
    last_table = doc.tables[-1]
    if len(last_table.rows) <= 4:
        rows_text = ' '.join(
            cell.text.strip()
            for row in last_table.rows[1:]
            for cell in row.cells
        )
        vague_terms = ['實體檔案銷毀程序', '刪除備援相關測試資料']
        if all(term in rows_text for term in vague_terms) and 'VM 申請表' not in rows_text \
                and '.vhd' not in rows_text.lower() and '虛擬硬碟' not in rows_text:
            findings.append((
                'RED',
                '捌、二、（五）備份或備援資料銷毀作業 Table',
                '銷毀程序過於空洞，沒寫具體方法（整機銷毀/硬碟刪除/VM 申請表佐證）'
            ))
    return findings


def rule_single_vs_dual_host_mismatch(doc):
    findings = []
    scenario_mentions_dual = False
    flow_mentions_dual = False

    for p in doc.paragraphs:
        if '演練情境' in p.text or '情境' in p.text:
            if ('兩台' in p.text) or ('伺服器與資料庫伺服器' in p.text) \
                    or ('WEB 主機與 DB 主機' in p.text):
                scenario_mentions_dual = True

    for t in doc.tables:
        for row in t.rows:
            full = ' '.join(cell.text for cell in row.cells)
            if ('DB 虛擬主機' in full) or ('AP 虛擬主機' in full):
                flow_mentions_dual = True
                break

    # Check if single-host architecture claimed
    single_host_claimed = False
    for p in doc.paragraphs:
        if '合併承載' in p.text or '單一主機合併' in p.text \
                or '一台備援虛擬主機' in p.text:
            single_host_claimed = True
            break

    if single_host_claimed and flow_mentions_dual:
        findings.append((
            'RED',
            'Table (緊急備用作業等)',
            '文字宣告「一台主機合併」但流程表仍提「DB 虛擬主機／AP 虛擬主機」兩台架構，不一致'
        ))

    return findings


def rule_extra_space_after_colon(doc):
    findings = []
    for i, p in enumerate(doc.paragraphs):
        # 冒號後多半形空格再接非空白
        if re.search(r'完成時間： +\d', p.text):
            findings.append((
                'GREEN',
                f'Paragraph [{i}]',
                '冒號後多餘半形空格'
            ))
    return findings


def rule_numbered_list_missing_dot(doc):
    findings = []
    for i, p in enumerate(doc.paragraphs):
        if re.match(r'^\d [A-Z]', p.text):
            findings.append((
                'GREEN',
                f'Paragraph [{i}]',
                f'項次編號缺點號：「{p.text[:20]}...」應為「1. ...」'
            ))
    return findings


ALL_RULES = [
    rule_unit_misplacement,
    rule_metrics_missing,
    rule_duplicate_phrase_in_purpose,
    rule_chinese_terms_typo,
    rule_empty_destruction,
    rule_single_vs_dual_host_mismatch,
    rule_extra_space_after_colon,
    rule_numbered_list_missing_dot,
]


def run_review(docx_path):
    doc = Document(docx_path)
    all_findings = []
    for rule in ALL_RULES:
        try:
            all_findings.extend(rule(doc))
        except Exception as e:
            all_findings.append(('GREEN', 'Rule error',
                                 f'{rule.__name__}: {e}'))

    by_severity = {'RED': [], 'YELLOW': [], 'GREEN': []}
    for sev, loc, msg in all_findings:
        by_severity[sev].append((loc, msg))

    return by_severity


def print_report(result, docx_path):
    red = result['RED']
    yel = result['YELLOW']
    grn = result['GREEN']

    print(f'Review: {docx_path}')
    print(f'🔴 {len(red)} 項必改 / 🟡 {len(yel)} 項建議改 / 🟢 {len(grn)} 項小瑕疵')
    print('=' * 60)

    if red:
        print('\n🔴 致命級（必改）')
        for i, (loc, msg) in enumerate(red, 1):
            print(f'{i}. [{loc}] {msg}')

    if yel:
        print('\n🟡 次級（建議改）')
        for i, (loc, msg) in enumerate(yel, 1):
            print(f'{i}. [{loc}] {msg}')

    if grn:
        print('\n🟢 小瑕疵')
        for i, (loc, msg) in enumerate(grn, 1):
            print(f'{i}. [{loc}] {msg}')

    if not (red or yel or grn):
        print('\n✅ 無發現問題')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    args = ap.parse_args()
    result = run_review(args.input)
    print_report(result, args.input)


if __name__ == '__main__':
    main()
