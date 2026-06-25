# -*- coding: utf-8 -*-
"""
BCP 演練計畫書 docx 產生器

輸入：JSON config
輸出：.docx（完整章節，符合 ISMS 規範）

用法：
    python gen_bcp_docx.py --config config.json --output ./output/XXX.docx

Config schema 參考 examples/example_config.json
"""
import argparse
import json
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class BCPDocBuilder:
    def __init__(self, config):
        self.cfg = config
        self.doc = Document()
        self._setup_style()

    def _setup_style(self):
        style = self.doc.styles['Normal']
        style.font.name = '標楷體'
        style.font.size = Pt(12)
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')

    # ---------- helpers ----------
    def _run(self, p, text, bold=False, size=12):
        r = p.add_run(text)
        r.font.name = '標楷體'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
        r.font.size = Pt(size)
        r.bold = bold
        return r

    def para(self, text='', bold=False, size=12, align=None, indent=0):
        p = self.doc.add_paragraph()
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        if text:
            self._run(p, text, bold=bold, size=size)
        return p

    def h(self, text, level=1):
        size = {1: 18, 2: 16, 3: 14, 4: 13}.get(level, 12)
        p = self.doc.add_paragraph()
        if level == 1:
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
        self._run(p, text, bold=True, size=size)
        return p

    def _set_cell_border(self, cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn('w:tcBorders'))
        if tcBorders is None:
            tcBorders = OxmlElement('w:tcBorders')
            tcPr.append(tcBorders)
        for edge in ('top', 'left', 'bottom', 'right'):
            tag = OxmlElement(f'w:{edge}')
            tag.set(qn('w:val'), 'single')
            tag.set(qn('w:sz'), '4')
            tag.set(qn('w:color'), '000000')
            existing = tcBorders.find(qn(f'w:{edge}'))
            if existing is not None:
                tcBorders.remove(existing)
            tcBorders.append(tag)

    def table(self, headers, rows, col_widths_cm=None):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(p, h, bold=True, size=11)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            self._set_cell_border(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'D9E1F2')
            tcPr.append(shd)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.text = ''
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._run(p, str(val), size=11)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                self._set_cell_border(cell)
        if col_widths_cm:
            for ri in range(len(t.rows)):
                for ci, w in enumerate(col_widths_cm):
                    t.rows[ri].cells[ci].width = Cm(w)
        return t

    def blank(self):
        self.doc.add_paragraph()

    def pagebreak(self):
        self.doc.add_page_break()

    def embed_image(self, path, width_cm=16, caption=None):
        if not path or not os.path.exists(path):
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(path, width=Cm(width_cm))
        if caption:
            c = self.doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(c, caption, bold=True, size=11)

    # ---------- sections ----------
    def build_cover(self):
        c = self.cfg
        self.para(c['agency'], bold=True, size=22, align='center')
        self.blank()
        self.para(c['system_name'], bold=True, size=22, align='center')
        self.blank()
        self.para('營運持續計畫', bold=True, size=22, align='center')
        self.blank(); self.blank(); self.blank()
        self.para(f"第 {c['version']} 版", bold=True, size=18, align='center')
        self.blank()
        self.para(c['drill_date_full'], bold=True, size=16, align='center')
        self.pagebreak()

    def build_toc(self):
        self.h('目  錄', level=1)
        self.para(
            '※ 請於 Word 開啟本檔後，點選「參考資料」→「目錄」更新目錄；'
            '或沿用既有目錄按 F9 更新頁碼。', size=10)
        self.pagebreak()

    def build_purpose(self):
        self.h('壹、目的', level=1)
        default_purpose = (
            f"為確保{self.cfg['agency']}（以下簡稱本署）之業務活動持續運作，"
            '不受重大災難、人為破壞或設備故障等中斷之影響，以保護重要業務及資訊運作，'
            '依照演練計畫排程表針對關鍵業務進行時機演練，並辦理資通安全通報演練，'
            '檢驗資安通報機制及應變能力。'
            '透過業務持續運作於事前研擬災害發生之情境，設計復原操作流程，'
            '以降低災害發生之衝擊；事中透過災害模擬演練之操作，驗證復原流程之有效性，'
            '並使相關人員熟稔操作程序，增加應變能力；事後檢討災害演練之缺失，'
            '作為精進災害演練計畫之參考依據。'
            '並實施備份還原演練，實作還原備份資料庫與資料接收服務至備援虛擬主機，'
            '作為佐證備份自動化成效與資通系統安全自主檢查之重要依據。'
        )
        self.para(self.cfg.get('purpose', default_purpose))

    def build_service_description(self):
        s = self.cfg.get('service_description', {})
        self.h('貳、系統服務說明', level=1)

        self.h('一、服務對象', level=2)
        self.para('本系統之主要服務對象包含：')
        for item in s.get('target_users', []):
            self.para(item, indent=1)

        self.h('二、服務重要性', level=2)
        self.para(s.get('importance', ''))

        self.h('三、中斷可能影響', level=2)
        self.para(s.get('impact_intro', '若本系統發生中斷，可能造成之影響如下：'))
        for item in s.get('impacts', []):
            self.para(item, indent=1)

    def build_architecture(self):
        a = self.cfg.get('architecture', {})
        self.h('參、系統架構', level=1)
        self.para('本系統架構如下圖所示：')
        if a.get('diagram_path'):
            self.embed_image(a['diagram_path'],
                             caption=a.get('diagram_caption', '圖 1、系統架構圖'))

        self.h('一、架構組成', level=2)
        self.para(a.get('composition_intro', '本系統整體架構由若干區塊組成：'))
        for block in a.get('blocks', []):
            self.para(block['title'], bold=True)
            if block.get('detail'):
                self.para(block['detail'], indent=1)

        self.h('二、資料流說明', level=2)
        self.para(a.get('data_flow', ''))

        self.h('三、架構圖與營運持續演練之對應', level=2)
        self.para(a.get('drill_mapping', ''))

    def build_scope(self):
        self.h('肆、範圍', level=1)
        self.para('本次演練參演單位（以下簡稱參演單位）為：')
        for u in self.cfg['units']:
            self.para(u, indent=1)
        self.blank()
        self.para('本次演練參演系統／服務（以下簡稱參演系統／服務）為：')
        for s in self.cfg['scope_services']:
            self.para(s, indent=1)

    def build_schedule(self):
        self.h('伍、時程', level=1)
        self.para(f"一、演練時間：{self.cfg['drill_date_full']}", indent=1)
        self.para(f"二、演練地點：{self.cfg['drill_location']}", indent=1)
        self.para(f"三、參演系統／服務：{self.cfg.get('scope_services_summary', self.cfg['system_name'])}", indent=1)

    def build_personnel(self):
        self.h('陸、參與人員與職責', level=1)
        rows = [
            [p['unit'], p['title'], p['name'], p['role']]
            for p in self.cfg['personnel']
        ]
        self.table(
            headers=['單位', '職稱', '姓名', '職責'],
            rows=rows,
            col_widths_cm=[4.5, 2.5, 2, 7.5]
        )

    def build_execution(self):
        e = self.cfg['execution']
        self.h('柒、執行方式', level=1)

        self.h('一、演練情境', level=3)
        self.para(e['scenario'])

        self.h('二、演練方式', level=3)
        self.para(e['method'])

        self.h('（一）通報與應變作業', level=4)
        self.para(e.get('response_desc',
            '由指揮中心依演練腳本逐項宣告事件進程，參演人員依「陸、參與人員與職責」所列角色'
            '執行通報、決策、資源調度與紀錄作業，驗證資通安全事件通報程序及緊急應變決策鏈之有效性。'
        ))

        for sec in e.get('hands_on_sections', []):
            self.h(sec['title'], level=4)
            if sec.get('intro'):
                self.para(sec['intro'])
            for i, step in enumerate(sec.get('steps', []), 1):
                self.para(f'{i}. {step}', indent=1)

        self.h('（五）整合驗證', level=4)
        for i, step in enumerate(e.get('verification_steps', []), 1):
            self.para(f'{i}. {step}', indent=1)

        self.h('（六）時間管控指標', level=4)
        m = e['metrics']
        self.para(f"1. MTPD（最大可容忍中斷時間）：{m['mtpd']}", indent=1)
        self.para(f"2. RTO（復原時間目標）：{m['rto']}", indent=1)
        self.para(f"3. RPO（復原點目標）：{m['rpo']}", indent=1)
        if m.get('worst_case_note'):
            self.blank()
            self.para('【本次演練情境說明】', bold=True)
            self.para(m['worst_case_note'])

        self.h('（七）備份與還原來源', level=4)
        for i, item in enumerate(e.get('backup_sources', []), 1):
            self.para(f'{i}. {item}', indent=1)

        self.h('（八）演練平台規格', level=4)
        for i, item in enumerate(e.get('platform_spec', []), 1):
            self.para(f'{i}. {item}', indent=1)

    def build_flow(self):
        f = self.cfg['flow']
        self.h('捌、執行作業流程', level=1)

        self.h('一、相關人員', level=2)
        self.h('（一）本署人員聯絡資訊', level=3)
        self.table(
            headers=['姓名', '職稱', '電話', '分機'],
            rows=[[c['name'], c['title'], c['phone'], c['ext']]
                  for c in self.cfg['contact_agency']],
            col_widths_cm=[3, 3, 4, 2.5]
        )
        self.h('（二）廠商聯絡資訊', level=3)
        self.table(
            headers=['姓名', '職稱', '電話', '分機'],
            rows=[[c['name'], c['title'], c['phone'], c['ext']]
                  for c in self.cfg['contact_vendor']],
            col_widths_cm=[3, 3.5, 4, 2.5]
        )

        self.h('二、緊急處理程序', level=2)

        for step_key, label in [
            ('incident_report', '（一）資通安全事件通報'),
            ('emergency_response', '（二）緊急應變作業'),
            ('backup_operation', '（三）緊急備用作業'),
            ('recovery', '（四）復原作業'),
        ]:
            step = f.get(step_key, {})
            self.h(label, level=3)
            if step.get('duration'):
                self.para(step['duration'])
            if step.get('rows'):
                self.table(
                    headers=['項次', '需求時間(min)', '採取行動', '負責人員／部門'],
                    rows=step['rows'],
                    col_widths_cm=[1.5, 2.5, 8, 4.5]
                )
            if step.get('note'):
                self.para(step['note'], size=10)

        # 銷毀作業
        dest = f.get('destruction', {})
        self.h('（五）備份或備援資料銷毀作業', level=3)
        if dest.get('rows'):
            self.table(
                headers=['項次', '採取行動', '負責人員／部門', '備註'],
                rows=dest['rows'],
                col_widths_cm=[1.5, 7.5, 3.5, 4]
            )
        if dest.get('note'):
            self.para(dest['note'], size=10)

    # ---------- orchestrate ----------
    def build_all(self):
        self.build_cover()
        self.build_toc()
        self.build_purpose()
        self.build_service_description()
        self.build_architecture()
        self.build_scope()
        self.build_schedule()
        self.build_personnel()
        self.build_execution()
        self.build_flow()

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--config', required=True, help='JSON config file')
    ap.add_argument('--output', required=True, help='Output docx path')
    args = ap.parse_args()

    with open(args.config, 'r', encoding='utf-8') as f:
        config = json.load(f)

    builder = BCPDocBuilder(config)
    builder.build_all()
    builder.save(args.output)
    size = os.path.getsize(args.output)
    print(f'OK: {args.output}  ({size:,} bytes)')


if __name__ == '__main__':
    main()
