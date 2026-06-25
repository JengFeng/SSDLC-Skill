"""
meeting-record / scripts / pack_pptx.py
========================================
把 slides/ 內的 PNG 包成一份 PPTX（每張 PNG = 一頁 full-bleed）。
給 proposal-narration pipeline 吃。

Usage:
    python pack_pptx.py [slides_dir] [out_pptx]

Default:
    slides/  →  video/meeting.pptx
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

sys.stdout.reconfigure(encoding='utf-8')


def pack(slides_dir: Path, out_pptx: Path):
    pngs = sorted(slides_dir.glob("*.png"))
    if not pngs:
        print(f"❌ No *.png found in {slides_dir}")
        return

    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for png in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(png), 0, 0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))
    print(f"✅ Packed {len(pngs)} slides → {out_pptx}")


if __name__ == "__main__":
    slides_dir = Path(sys.argv[1] if len(sys.argv) > 1 else "slides")
    out_pptx = Path(sys.argv[2] if len(sys.argv) > 2 else "video/meeting.pptx")
    pack(slides_dir, out_pptx)
