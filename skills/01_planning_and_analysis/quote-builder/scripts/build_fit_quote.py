#!/usr/bin/env python
"""
FIT Quote Builder — 產生「準線智慧科技」公司專用報價單 .xlsx

照 FT26050006 版型從零以 openpyxl 建構，不依賴 .xlsx 範本檔。
與 build_quote.py（政府案經費明細表格式）並存，互不影響。

版型重點：
  公司抬頭 → 左右雙欄基本資料 → 表格(項次/項目及說明/數量/單位/單價/金額)
  → 小計 → 稅額5% → 總計 → 5 點備註 → 客戶用印回傳區

用法 1（CLI + JSON config）：
  python build_fit_quote.py --output out.xlsx --config quote.json

用法 2（import）：
  from build_fit_quote import build_fit_quote
  build_fit_quote(output, header, groups, notes=None, tax_pct=0.05)

config / 參數結構見 references/fit_format.md
"""
import argparse, json, math, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ── 公司固定資訊（依 FT26050006）──────────────────────────────
COMPANY = {
    "name_zh":  "準線智慧科技",
    "name_en":  "Focus Intelligence Technology",
    "address":  "414 臺中市烏日區高鐵一路 268 號 18 樓之 6",
    "tel_fax":  "TEL：04-2451-6609　FAX：04-2463-1362",
    "url":      "https://www.focusit.com.tw/",
}

# FIT 制式 5 點備註；caller 傳 notes 可整組覆寫
DEFAULT_NOTES = [
    "報價單有效期：30 天內有效。本報價單經客戶簽回且加蓋公司章視同確認訂單，本公司得保留接受本訂單之權利。",
    "本報價不含備品、耗材、硬體設備採購及故障設備更換費用。若執行期間發生設備異常，將由本公司先行檢測判斷；經判定為系統異常者，由本公司負責進行故障排除與系統恢復作業(不另行收費)。若涉及現場硬體設備損壞、零組件故障、耗材更換或設備更換需求，將另行提出處理建議及報價，經業主確認後再辦理採購、維修或更換作業。",
    "緊急維修支援費單次出勤費用為新臺幣 12,000 元。後續如有現場維護、緊急維修或現場支援需求，將依本項費用標準另行報價收費。零組件、備品、耗材及設備更換費用另計。",
    "如有特殊發票開立需求（包含分開開立、抬頭異動等），應於回簽前提出並完成確認。報價單一經回簽，視同同意本公司發票開立規範，恕不受理後續任何變更申請。",
    "本報價單內容非經雙方協議修改者，該修改視為無效。",
]

_ZH = ["", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
       "十一", "十二", "十三", "十四", "十五"]

FN = "微軟正黑體"


def _vlines(lines, cpl):
    """估算多行文字在指定欄寬下的視覺行數（中文字寬約 latin 2 倍）。"""
    n = 0
    for ln in lines:
        n += max(1, math.ceil(len(ln) / cpl))
    return n


def build_fit_quote(output, header, groups, notes=None, tax_pct=0.05, company=None):
    """
    header: dict，客戶/報價單抬頭欄位（見 references/fit_format.md）
    groups: list of dict，每組 {"title": 大項名稱(可空字串=不分組), "items": [...]}
            item: {"name", "desc"(list或str,可空), "qty", "unit", "price"}
    notes:  list[str]，None 則用 FIT 制式 5 點
    回傳：dict {小計, 稅, 總計}
    """
    company = {**COMPANY, **(company or {})}
    notes = list(notes) if notes is not None else DEFAULT_NOTES

    wb = Workbook()
    ws = wb.active
    ws.title = "報價單"

    widths = {"A": 6, "B": 52, "C": 8, "D": 10, "E": 13, "F": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    gray = PatternFill("solid", fgColor="D9D9D9")

    f_base   = Font(name=FN, size=10)
    f_bold   = Font(name=FN, size=10, bold=True)
    f_label  = Font(name=FN, size=10, bold=True)
    f_title  = Font(name=FN, size=14, bold=True)
    f_quote  = Font(name=FN, size=22, bold=True)
    f_quoteE = Font(name=FN, size=11, italic=True)
    f_grand  = Font(name=FN, size=11, bold=True)

    a_lc  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    a_l   = Alignment(horizontal="left",   vertical="center")
    a_lt  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
    a_c   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    a_r   = Alignment(horizontal="right",  vertical="center")

    def put(coord, value, *, font=f_base, align=None, border=None, fill=None, numfmt=None):
        cell = ws[coord]
        cell.value = value
        cell.font = font
        if align:  cell.alignment = align
        if border: cell.border = border
        if fill:   cell.fill = fill
        if numfmt: cell.number_format = numfmt
        return cell

    def border_range(rng):
        for row in ws[rng]:
            for cell in row:
                cell.border = box

    def fill_range(rng):
        for row in ws[rng]:
            for cell in row:
                cell.fill = gray

    r = 1
    # ── 公司抬頭 ───────────────────────────────────────────
    ws.merge_cells(f"A{r}:C{r}")
    put(f"A{r}", f"{company['name_zh']}　{company['name_en']}", font=f_title, align=a_l)
    ws.merge_cells(f"D{r}:F{r+1}")
    put(f"D{r}", "報　價　單", font=f_quote, align=a_c)
    r += 1
    ws.merge_cells(f"A{r}:C{r}")
    put(f"A{r}", company["address"], align=a_l)
    r += 1
    ws.merge_cells(f"A{r}:C{r}")
    put(f"A{r}", company["tel_fax"], align=a_l)
    ws.merge_cells(f"D{r}:F{r}")
    put(f"D{r}", "Quotation", font=f_quoteE, align=a_c)
    r += 1
    ws.merge_cells(f"A{r}:C{r}")
    put(f"A{r}", company["url"], align=a_l)
    r += 2  # 空白列

    # ── 左右雙欄基本資料 ───────────────────────────────────
    left_fields = [
        ("客戶名稱", header.get("client_name", "")),
        ("客戶統編", header.get("client_tax_id", "")),
        ("聯絡人",   header.get("contact", "")),
        ("聯絡電話", header.get("contact_phone", "")),
        ("發票地址", header.get("invoice_address", "")),
        ("專案名稱", header.get("project", "")),
    ]
    right_fields = [
        ("報價單號",   header.get("quote_no", "")),
        ("報價日期",   header.get("quote_date", "")),
        ("報價有效期", header.get("valid_until", "")),
        ("業務聯絡人", header.get("sales", "")),
        ("分機",       header.get("ext", "")),
        ("電子信箱",   header.get("email", "")),
    ]
    fstart = r
    for i in range(6):
        rr = fstart + i
        lab_l, val_l = left_fields[i]
        put(f"A{rr}", lab_l, font=f_label, align=a_l)
        ws.merge_cells(f"B{rr}:C{rr}")
        put(f"B{rr}", str(val_l), align=a_lc)
        lab_r, val_r = right_fields[i]
        put(f"D{rr}", lab_r, font=f_label, align=a_l)
        ws.merge_cells(f"E{rr}:F{rr}")
        put(f"E{rr}", str(val_r), align=a_lc)
    r = fstart + 6 + 1  # 空白列

    # ── 表頭 ──────────────────────────────────────────────
    for i, h in enumerate(["項次", "項目及說明", "數量", "單位", "單價", "金額"]):
        put(f"{get_column_letter(i+1)}{r}", h, font=f_bold, align=a_c, border=box, fill=gray)
    ws.row_dimensions[r].height = 24
    r += 1

    # ── 工項列（可分大項組）────────────────────────────────
    item_rows = []
    item_no = 0
    gi = 0
    for g in groups:
        title = (g.get("title") or "").strip()
        if title:
            gi += 1
            ws.merge_cells(f"A{r}:F{r}")
            num = _ZH[gi] if gi < len(_ZH) else str(gi)
            put(f"A{r}", f"{num}、{title}", font=f_bold, align=a_l)
            border_range(f"A{r}:F{r}")
            fill_range(f"A{r}:F{r}")
            ws.row_dimensions[r].height = 22
            r += 1
        for it in g.get("items", []):
            item_no += 1
            desc = it.get("desc", [])
            if isinstance(desc, str):
                desc = [desc] if desc else []
            lines = [it["name"]] + [f"•{d}" for d in desc]
            put(f"A{r}", item_no, align=a_c, border=box)
            put(f"B{r}", "\n".join(lines), align=a_lt, border=box)
            put(f"C{r}", it.get("qty", 1), align=a_c, border=box)
            put(f"D{r}", it.get("unit", "式"), align=a_c, border=box)
            put(f"E{r}", it.get("price", 0), align=a_r, border=box, numfmt="#,##0")
            put(f"F{r}", f"=C{r}*E{r}", align=a_r, border=box, numfmt="#,##0")
            ws.row_dimensions[r].height = 17 * _vlines(lines, 24) + 8
            item_rows.append(r)
            r += 1

    # ── 小計 / 稅額 / 總計 ─────────────────────────────────
    sub_r = r
    if item_rows:
        sum_f = f"=SUM(F{item_rows[0]}:F{item_rows[-1]})"
    else:
        sum_f = 0
    rows = [
        ("小計(Subtotal)",                f_bold,  sum_f),
        (f"稅額 {int(round(tax_pct*100))}%(Tax)", f_bold,  f"=F{sub_r}*{tax_pct}"),
        ("總計(Total)",                   f_grand, f"=F{sub_r}+F{sub_r+1}"),
    ]
    for i, (lab, lab_font, val) in enumerate(rows):
        rr = sub_r + i
        ws.merge_cells(f"A{rr}:E{rr}")
        put(f"A{rr}", lab, font=lab_font, align=a_r, border=box)
        border_range(f"A{rr}:E{rr}")
        put(f"F{rr}", val, font=(f_grand if i == 2 else f_base),
            align=a_r, border=box, numfmt="#,##0")
    r = sub_r + 3 + 1  # 空白列

    # ── 備註 ──────────────────────────────────────────────
    put(f"A{r}", "備註：", font=f_bold, align=a_l)
    r += 1
    for i, note in enumerate(notes, start=1):
        ws.merge_cells(f"A{r}:F{r}")
        put(f"A{r}", f"{i}. {note}", align=a_lt)
        ws.row_dimensions[r].height = 17 * _vlines([note], 48) + 8
        r += 1
    r += 1  # 空白列

    # ── 客戶用印回傳 ──────────────────────────────────────
    ws.merge_cells(f"A{r}:F{r}")
    put(f"A{r}", "客戶確認用印回傳", font=f_bold, align=a_l)
    r += 1
    put(f"A{r}", "發票章或公司章", align=a_l)
    put(f"D{r}", "簽回日期：", align=a_l)
    ws.row_dimensions[r].height = 48

    # ── 版面 ──────────────────────────────────────────────
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = f"A1:F{r}"

    wb.save(output)

    sub = sum(it.get("qty", 1) * it.get("price", 0)
              for g in groups for it in g.get("items", []))
    tax = sub * tax_pct
    return {"小計": sub, "稅": tax, "總計": sub + tax}


def main():
    p = argparse.ArgumentParser(description="Build .xlsx quote in FIT 公司專用格式")
    p.add_argument("--output", "-o", required=True, help="輸出 .xlsx 路徑")
    p.add_argument("--config", required=True, help="JSON config 檔路徑")
    args = p.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    result = build_fit_quote(
        args.output,
        header=cfg["header"],
        groups=cfg["groups"],
        notes=cfg.get("notes"),
        tax_pct=cfg.get("tax_pct", 0.05),
        company=cfg.get("company"),
    )
    print(f"輸出檔：{args.output}")
    print(f"小計   = {result['小計']:,.0f}")
    print(f"稅額   = {result['稅']:,.0f}")
    print(f"總計   = {result['總計']:,.0f}")


if __name__ == "__main__":
    main()
